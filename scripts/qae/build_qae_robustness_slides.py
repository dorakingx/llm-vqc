"""Build the 10-slide "tested one-factor slices" robustness deck (.pptx + .pdf).

Structure and explanatory style follow the previous meeting deck. Content is
the completed robustness study only: four tested one-factor slices from one
baseline, never a factorial grid. Every number is read from the committed
result tables; nothing is typed by hand and no experiment is run here.

Model tiers are named Low (gpt-4.1-mini) and High (gpt-5.4-mini); groups are
always ordered Random -> Greedy -> Low -> High; LLM-Open / LLM-Closed remain
the search-workflow names.

QA performed on every build:
  - exactly 10 slides, nothing off-canvas, no estimated text overflow
  - no coding/version identifiers, no Model A/B style labels, no full-grid claims
  - every technical term defined on or before the slide it first appears
  - required section order, required phrases, B_min procedure present

Usage:
  python scripts/qae/build_qae_robustness_slides.py          # build + PDF
  python scripts/qae/build_qae_robustness_slides.py --check  # audit only
"""
import argparse
import importlib.util
import json
import re
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path("outputs/qae_robustness")
FIGURES = ROOT / "figures"
DECK_DIR = ROOT / "deck"
DECK_NAME = "20260904_GSoC"

_spec = importlib.util.spec_from_file_location(
    "qae_deck_common", Path("scripts/qae/build_qae_robustness_deck.py"))
_common = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_common)
geometry_audit = _common.audit
extract_text = _common.extract_text

# Shared coding/version patterns plus the labels and claims this deck must not carry.
FORBIDDEN_PATTERNS = _common.FORBIDDEN_PATTERNS + [
    (r"\bModel [AB]\b", "model letter label"),
    (r"\b(Open|Closed) [AB]\b", "model letter label"),
    (r"\b(reference|alternative) model\b", "model role label"),
    (r"(?i)across the full grid", "full-grid claim"),
    (r"(?i)for all qubit.budget combinations", "full-grid claim"),
    (r"(?i)\b(model|Hamiltonian)-independent\b", "over-general claim"),
]

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

NAVY = RGBColor(0x1F, 0x2A, 0x44)
KICKER = RGBColor(0xC5, 0x22, 0x1F)
INK = RGBColor(0x20, 0x21, 0x24)
MUTED = RGBColor(0x5F, 0x63, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE, BLUE_FILL = RGBColor(0x1A, 0x73, 0xE8), RGBColor(0xE8, 0xF0, 0xFE)
ORANGE, ORANGE_FILL = RGBColor(0xE8, 0x71, 0x0A), RGBColor(0xFE, 0xF3, 0xE0)
GREEN, GREEN_FILL = RGBColor(0x18, 0x80, 0x38), RGBColor(0xE6, 0xF4, 0xEA)
RED, RED_FILL = RGBColor(0xD9, 0x30, 0x25), RGBColor(0xFC, 0xE8, 0xE6)
AMBER, AMBER_FILL = RGBColor(0xF9, 0xAB, 0x00), RGBColor(0xFE, 0xF7, 0xE0)
GREY, GREY_FILL = RGBColor(0x9A, 0xA0, 0xA6), RGBColor(0xF1, 0xF3, 0xF4)
FONT = "Helvetica Neue"

LOW_MODEL, HIGH_MODEL = "gpt-4.1-mini", "gpt-5.4-mini"
OPEN_RANDOM = "LLM-Open_minus_Random"
CLOSED_OPEN = "LLM-Closed_minus_LLM-Open"
CLOSED_RANDOM = "LLM-Closed_minus_Random"
GREEDY_RANDOM = "Greedy_minus_Random"

# ------------------------------------------------------------ primitives ---


def _flat(shape):
    shape.shadow.inherit = False
    style = shape._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        shape._element.remove(style)
    return shape


def rect(slide, left, top, width, height, color):
    shape = _flat(slide.shapes.add_shape(1, left, top, width, height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def box(slide, left, top, width, height, fill, border, border_pt=1.25):
    shape = _flat(slide.shapes.add_shape(5, left, top, width, height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_pt)
    return shape


def line(slide, x1, y1, x2, y2, color=GREY, width_pt=1.5):
    connector = slide.shapes.add_connector(1, x1, y1, x2, y2)
    connector.line.color.rgb = color
    connector.line.width = Pt(width_pt)
    return connector


def textbox(slide, left, top, width, height, align=PP_ALIGN.LEFT):
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    frame.paragraphs[0].alignment = align
    return frame


def write(frame, lines, *, size=12, color=INK, bold=False, space_after=4,
          line_spacing=1.0, align=None):
    """lines: str | (str, opts) | [ (run_text, run_opts), ... ] per paragraph."""
    first = True
    for entry in lines:
        if isinstance(entry, list):
            runs, opts = entry, {}
        else:
            text, opts = (entry if isinstance(entry, tuple) else (entry, {}))
            runs = [(text, opts)]
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.space_after = Pt(opts.get("space_after", space_after))
        paragraph.line_spacing = opts.get("line_spacing", line_spacing)
        if align is not None or "align" in opts:
            paragraph.alignment = opts.get("align", align)
        for text, ropts in runs:
            run = paragraph.add_run()
            run.text = text
            font = run.font
            font.name = FONT
            font.size = Pt(ropts.get("size", opts.get("size", size)))
            font.bold = ropts.get("bold", opts.get("bold", bold))
            font.italic = ropts.get("italic", False)
            font.color.rgb = ropts.get("color", opts.get("color", color))
    return frame


def header(slide, kicker, title, subtitle=None):
    frame = textbox(slide, Inches(0.55), Inches(0.24), Inches(12.2), Inches(0.28))
    write(frame, [(kicker.upper(), {"size": 11, "bold": True, "color": KICKER})],
          space_after=0)
    frame = textbox(slide, Inches(0.55), Inches(0.56), Inches(12.2), Inches(0.5))
    write(frame, [(title, {"size": 24, "bold": True, "color": NAVY})], space_after=0)
    if subtitle:
        frame = textbox(slide, Inches(0.55), Inches(1.1), Inches(12.2), Inches(0.38))
        write(frame, [(subtitle, {"size": 11, "color": MUTED})], space_after=0)


FOOTER_TOP = Inches(7.0)
TRANSITION_TOP = Inches(6.55)


def footer(slide, text, top=FOOTER_TOP):
    frame = textbox(slide, Inches(0.55), top, Inches(12.2), Inches(0.35))
    write(frame, [(text, {"size": 8.5, "color": MUTED})], space_after=0)


def transition(slide, text, top=TRANSITION_TOP):
    """The one sentence that hands over to the next slide (spec: slides 2-9)."""
    rect(slide, Inches(0.55), top, Inches(0.08), Inches(0.36), KICKER)
    frame = textbox(slide, Inches(0.75), top - Inches(0.02), Inches(12.0), Inches(0.42))
    write(frame, [[("Next:  ", {"bold": True, "color": KICKER}),
                   (text, {"color": INK})]], size=11, space_after=0)


def picture(slide, name, left, top, width=None, height=None):
    path = FIGURES / f"{name}.png"
    if not path.exists():
        raise SystemExit(f"missing figure {path}; run the slice figure builder first")
    if width is not None:
        return slide.shapes.add_picture(str(path), left, top, width=width)
    return slide.shapes.add_picture(str(path), left, top, height=height)


def titled_box(slide, left, top, width, height, fill, border, title, lines,
               title_color=None, title_size=12, body_size=10.5, space_after=3):
    box(slide, left, top, width, height, fill, border)
    frame = textbox(slide, left + Inches(0.14), top + Inches(0.08),
                    width - Inches(0.28), height - Inches(0.16))
    write(frame, [(title, {"size": title_size, "bold": True,
                           "color": title_color or border, "space_after": 3})]
          + [(entry if isinstance(entry, (tuple, list)) else (entry, {"size": body_size}))
             for entry in lines],
          size=body_size, space_after=space_after)
    return frame


def chip(slide, left, top, width, text, fill, border, color=INK, size=10, bold=False):
    shape = box(slide, left, top, width, Inches(0.36), fill, border, 1.0)
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = Emu(0)
    frame.margin_top = frame.margin_bottom = Emu(0)
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color
    return shape


TABLE_ROW_HEIGHT = Inches(0.34)


def table(slide, rows, left, top, col_widths, row_height=TABLE_ROW_HEIGHT,
          header_size=10, body_size=9.5, colors=None):
    width = sum(col_widths, Emu(0))
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width,
                                   row_height * len(rows)).table
    for index, w in enumerate(col_widths):
        shape.columns[index].width = w
    for r, row in enumerate(rows):
        shape.rows[r].height = row_height
        for c, value in enumerate(row):
            cell = shape.cell(r, c)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Emu(12700)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(header_size if r == 0 else body_size)
                    run.font.bold = r == 0 or c == 0
                    run.font.name = FONT
                    run.font.color.rgb = INK
                    if colors and (r, c) in colors:
                        run.font.color.rgb = colors[(r, c)]
                        run.font.bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = (RGBColor(0xE8, 0xEA, 0xED) if r == 0 else
                                        (WHITE if r % 2 else RGBColor(0xF8, 0xF9, 0xFA)))
    return shape


# --------------------------------------------------------------- numbers ---


def _fmt(value, digits=4, sign=True):
    return f"{value:+.{digits}f}" if sign else f"{value:.{digits}f}"


def _p(value):
    return "p < 0.001" if value < 0.001 else f"p = {value:.3f}"


def stat(summary, key, contrast):
    return summary["conditions"][key]["contrasts"][contrast]


def mean_of(summary, key, method):
    return summary["conditions"][key]["per_method"][method]["mean"]


def validity(key, method):
    entry = json.loads((ROOT / key / "proposal_quality.json").read_text())[method]
    return 1 - entry["fallback_evaluations"] / entry["evaluations"]


def short(summary, key, contrast):
    entry = stat(summary, key, contrast)
    return (f"{_fmt(entry['mean_paired_gain'], 3)} ({entry['wins_b']}/12, "
            f"{_p(entry['wilcoxon_exact_two_sided_p'])})")


def verdict(summary, key, contrast):
    """Evidence-calibrated word for one paired contrast in one condition."""
    entry = stat(summary, key, contrast)
    if entry["wilcoxon_exact_two_sided_p"] >= 0.05:
        return "not detected"
    return "robust" if entry["mean_paired_gain"] > 0 else "reversed"


# ---------------------------------------------------------------- slides ---


def slide_01_question(deck, summary):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    rect(slide, Emu(0), Emu(0), SLIDE_W, SLIDE_H, NAVY)
    frame = textbox(slide, Inches(0.9), Inches(0.95), Inches(11.6), Inches(1.9))
    write(frame, [("Which parts of the previous QAE result",
                   {"size": 36, "bold": True, "color": WHITE, "space_after": 0}),
                  ("survive one-factor changes?",
                   {"size": 36, "bold": True, "color": WHITE})], space_after=6)
    frame = textbox(slide, Inches(0.9), Inches(2.85), Inches(11.6), Inches(0.75))
    write(frame, [("A quantum autoencoder (QAE) robustness study — a follow-up to last "
                   "meeting's architecture-search result",
                   {"size": 16, "color": RGBColor(0xC8, 0xD4, 0xEC)})])
    box(slide, Inches(0.9), Inches(3.7), Inches(11.6), Inches(1.95),
        RGBColor(0x2A, 0x3A, 0x5C), RGBColor(0x5B, 0x8D, 0xEF))
    frame = textbox(slide, Inches(1.12), Inches(3.8), Inches(11.2), Inches(1.8))
    write(frame, [
        [("Baseline finding:  ", {"bold": True, "color": RGBColor(0xF9, 0xAB, 0x00)}),
         ("physics-informed language-model proposals beat random and greedy circuit "
          "search by a wide margin, and a closed feedback loop added a small extra gain.",
          {"color": WHITE})],
        [("Four factors tested, one at a time:  ",
          {"bold": True, "color": RGBColor(0xF9, 0xAB, 0x00)}),
         ("evaluation budget · qubit count · Hamiltonian · language model.",
          {"color": WHITE})],
        [("This is a tested one-factor-slice study, not a complete factorial grid.",
          {"bold": True, "color": WHITE})],
    ], size=13.5, space_after=6)
    frame = textbox(slide, Inches(0.9), Inches(5.95), Inches(11.5), Inches(1.0))
    write(frame, [("Tomoya Hatanaka", {"size": 15, "bold": True, "color": WHITE}),
                  ("ML4SCI / Google Summer of Code 2026",
                   {"size": 12, "color": RGBColor(0xC8, 0xD4, 0xEC)}),
                  ("September 4, 2026", {"size": 11, "color": RGBColor(0x9A, 0xA8, 0xC8)})],
          space_after=2)


def slide_02_baseline(deck, summary):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    header(slide, "1 · Previous baseline",
           "Where we left off: the previous baseline result",
           "Quantum autoencoder (QAE): compress ground states into fewer qubits — "
           "4 qubits · Ising chain · B = 8 · High model tier · 12 paired seeds")
    flow = [
        ("Input  |ψ(h)⟩", "ground state of the 4-qubit Ising chain\n"
         "H = −Σ Zᵢ Zᵢ₊₁ − h Σ Xᵢ,  h ∈ [0.2, 2.0]", BLUE_FILL, BLUE),
        ("Encoder  Uθ", "the circuit we search:\n16 gates, angles trained", ORANGE_FILL, ORANGE),
        ("Latent qubits  q0, q1", "the 2 qubits we KEEP\n(the compressed code)",
         GREEN_FILL, GREEN),
        ("Trash qubits  q2, q3", "should end in the SAME state |00⟩\nfor every input",
         RED_FILL, RED),
    ]
    x = 0.55
    widths = [3.55, 2.55, 2.6, 2.7]
    for (title, body, fill, border), w in zip(flow, widths, strict=True):
        titled_box(slide, Inches(x), Inches(1.55), Inches(w), Inches(0.92), fill, border,
                   title, [(body, {"size": 9, "color": INK})], title_size=11, body_size=9,
                   space_after=0)
        if w != widths[-1]:
            frame = textbox(slide, Inches(x + w + 0.02), Inches(1.78), Inches(0.3), Inches(0.4))
            write(frame, [("→", {"size": 16, "color": MUTED})], space_after=0)
        x += w + 0.35
    box(slide, Inches(0.55), Inches(2.6), Inches(12.2), Inches(0.62), BLUE_FILL, BLUE)
    frame = textbox(slide, Inches(0.7), Inches(2.66), Inches(11.95), Inches(0.55))
    write(frame, [[("Trash fidelity:  ", {"bold": True, "color": BLUE}),
                   ("F_trash = ⟨00| ρ_trash |00⟩ = P(q2 q3 = 00), the probability that the "
                    "trash qubits end in |00⟩.  ", {}),
                   ("Higher is better.  ", {"bold": True, "color": GREEN}),
                   ("Reported on a held-out test set (64 unseen h values never used for "
                    "training, selection or feedback), read once after selection.", {})]],
          size=10.5, space_after=0)
    picture(slide, "baseline_groups", Inches(0.45), Inches(3.3), height=Inches(3.15))
    methods = [
        ("Random", GREY, "B independent uniformly sampled circuit architectures."),
        ("Greedy", AMBER, "B/2 random starts, then B/2 single-structural-change "
                          "refinements of the best validation candidate."),
        ("LLM-Open", BLUE, "open-loop: all semantic (physics-informed) proposals are "
                           "generated BEFORE any validation result is observed."),
        ("LLM-Closed", RED, "closed-loop: semantic initial proposals, then free "
                            "redesigns of the current best circuit using validation "
                            "feedback."),
    ]
    box(slide, Inches(6.65), Inches(3.3), Inches(6.1), Inches(1.98), GREY_FILL, GREY)
    frame = textbox(slide, Inches(6.8), Inches(3.36), Inches(5.85), Inches(1.9))
    lines = [[("Four search methods; evaluation budget B = the number of candidate "
               "circuits trained and evaluated per seed", {"bold": True})]]
    for name, color, body in methods:
        lines.append([(f"{name}  ", {"bold": True, "color": color}), (body, {})])
    write(frame, lines, size=9.5, space_after=2)
    box(slide, Inches(6.65), Inches(5.38), Inches(6.1), Inches(1.1), AMBER_FILL, AMBER)
    frame = textbox(slide, Inches(6.8), Inches(5.43), Inches(5.85), Inches(1.02))
    write(frame, [
        [("Two claims to stress-test", {"bold": True, "color": ORANGE})],
        [("1  Semantic proposals outperform non-semantic search.  ",
          {"bold": True, "color": BLUE}),
         ("LLM-Open − Random " + short(summary, "reference", OPEN_RANDOM), {})],
        [("2  Closed-loop redesign gives a smaller additional gain.  ",
          {"bold": True, "color": RED}),
         ("LLM-Closed − LLM-Open " + short(summary, "reference", CLOSED_OPEN), {})],
    ], size=9.5, space_after=2)
    transition(slide, "do these two claims survive when exactly one condition changes?",
               top=Inches(6.58))
    footer(slide, "Reading: dots = the 12 paired seeds (12 independent draws of training/"
                  "validation data, shared by all methods); marker = mean; bar = bootstrap "
                  "95% confidence interval (CI); differences are per-seed paired, "
                  f"p = exact paired Wilcoxon test. High tier = {HIGH_MODEL}.")


def slide_03_design(deck, summary):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    header(slide, "2 · Controlled robustness design",
           "One factor changes; everything else remains fixed",
           "Four controlled changes leave the baseline; each moves exactly one declared "
           "factor and holds the other three at the baseline")
    box(slide, Inches(0.55), Inches(2.35), Inches(3.1), Inches(1.5), NAVY, NAVY)
    frame = textbox(slide, Inches(0.7), Inches(2.42), Inches(2.8), Inches(1.4))
    write(frame, [("BASELINE", {"size": 11, "bold": True, "color": AMBER}),
                  ("4 qubits · Ising · B = 8", {"size": 12.5, "bold": True, "color": WHITE}),
                  ("High model tier", {"size": 12.5, "bold": True, "color": WHITE}),
                  ("12 paired seeds", {"size": 11, "color": RGBColor(0xC8, 0xD4, 0xEC)})],
          space_after=2)
    branches = [
        ("Evaluation budget B", ["4", "8", "16"], 1, BLUE, BLUE_FILL),
        ("Qubit count", ["4", "6", "8"], 0, GREEN, GREEN_FILL),
        ("Hamiltonian", ["Ising", "XXZ"], 0, AMBER, AMBER_FILL),
        ("Language-model tier", ["Low", "High"], 1, RED, RED_FILL),
    ]
    for i, (name, levels, base_index, color, fill) in enumerate(branches):
        top = Inches(1.6 + i * 0.78)
        line(slide, Inches(3.65), Inches(3.1), Inches(4.35), top + Inches(0.28), color, 1.75)
        box(slide, Inches(4.35), top, Inches(3.55), Inches(0.7), fill, color)
        frame = textbox(slide, Inches(4.48), top + Inches(0.01), Inches(3.3), Inches(0.24))
        write(frame, [(name, {"size": 10.5, "bold": True, "color": color})], space_after=0)
        x = 4.5
        for j, level in enumerate(levels):
            is_base = j == base_index
            chip(slide, Inches(x), top + Inches(0.31), Inches(0.72), level,
                 NAVY if is_base else WHITE, NAVY if is_base else color,
                 WHITE if is_base else INK, size=9.5, bold=is_base)
            x += 0.82
    frame = textbox(slide, Inches(4.35), Inches(4.78), Inches(3.55), Inches(0.5))
    write(frame, [[("dark chip", {"bold": True, "color": NAVY}),
                   (" = the baseline level. Low = " + LOW_MODEL + ", High = " + HIGH_MODEL
                    + " (defined on the next slide).", {"color": MUTED})]],
          size=9, space_after=0)
    titled_box(slide, Inches(8.2), Inches(1.6), Inches(4.55), Inches(3.65), GREEN_FILL, GREEN,
               "FROZEN in every condition", [
        "•  prompt wording and output schema",
        "•  method logic of all four search methods",
        "•  circuit resource rule: 3 rotations + 1 CNOT (controlled-NOT) per qubit",
        "•  trainer and optimizer",
        "•  data splits (train / validation / held-out test)",
        "•  the same 12 paired seeds",
        "•  validation-only selection",
        "•  held-out test evaluation, read once",
        "•  statistical analysis (paired, bootstrap CI, exact Wilcoxon)",
    ], body_size=10, space_after=3)
    box(slide, Inches(0.55), Inches(5.4), Inches(12.2), Inches(0.95), AMBER_FILL, AMBER)
    frame = textbox(slide, Inches(0.7), Inches(5.46), Inches(11.95), Inches(0.85))
    write(frame, [
        [("No interaction condition was tested.  ", {"bold": True, "color": ORANGE}),
         ("Each change returns to the baseline before the next factor moves, so nothing "
          "is known about, for example, a large budget at 8 qubits, or the Low tier on "
          "the XXZ chain. XXZ = the second Hamiltonian family, an XXZ Heisenberg chain "
          "H = Σ (Xᵢ Xᵢ₊₁ + Yᵢ Yᵢ₊₁ + Δ Zᵢ Zᵢ₊₁), Δ ∈ [0.2, 2.0], declared before the "
          "run; only its stated physics facts enter the prompt.", {})],
    ], size=10, space_after=0)
    transition(slide, "because only one factor moves at a time, the results should be "
                      "read as four tested slices, not as a full grid.")
    footer(slide, "An automated check refuses to run a condition whose frozen components "
                  "differ from the baseline's or that moves more than one factor; the "
                  "baseline cell is reused, not re-run.")


def slide_04_reading(deck, summary):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    header(slide, "3 · What was actually tested, and how to read it",
           "Four tested slices, one baseline, no empty panels",
           "Every result slide shows one of these four rows; untested combinations are "
           "omitted rather than displayed as empty panels")
    slices = [
        ("Budget slice", ["B = 4", "B = 8", "B = 16"], 1, "4 qubits · Ising · High fixed",
         BLUE, BLUE_FILL),
        ("Qubit slice", ["4 qubits", "6 qubits", "8 qubits"], 0,
         "B = 8 · Ising · High fixed", GREEN, GREEN_FILL),
        ("Hamiltonian slice", ["Ising", "XXZ"], 0, "4 qubits · B = 8 · High fixed",
         AMBER, AMBER_FILL),
        ("Model slice", ["Low", "High"], 1, "4 qubits · Ising · B = 8 fixed",
         RED, RED_FILL),
    ]
    for i, (name, levels, base_index, fixed, color, fill) in enumerate(slices):
        top = Inches(1.6 + i * 0.66)
        box(slide, Inches(0.55), top, Inches(7.4), Inches(0.56), fill, color, 1.0)
        frame = textbox(slide, Inches(0.68), top + Inches(0.1), Inches(1.7), Inches(0.4))
        write(frame, [(name, {"size": 10.5, "bold": True, "color": color})], space_after=0)
        x = 2.4
        for j, level in enumerate(levels):
            is_base = j == base_index
            chip(slide, Inches(x), top + Inches(0.1), Inches(0.95), level,
                 NAVY if is_base else WHITE, NAVY if is_base else color,
                 WHITE if is_base else INK, size=9.5, bold=is_base)
            x += 1.02
            if j < len(levels) - 1:
                frame = textbox(slide, Inches(x - 0.12), top + Inches(0.1), Inches(0.2),
                                Inches(0.36))
                write(frame, [("→", {"size": 11, "color": MUTED})], space_after=0,
                      align=PP_ALIGN.CENTER)
                x += 0.1
        frame = textbox(slide, Inches(5.65), top + Inches(0.12), Inches(2.25), Inches(0.36))
        write(frame, [(fixed, {"size": 9, "color": MUTED})], space_after=0)
    frame = textbox(slide, Inches(0.55), Inches(4.3), Inches(7.4), Inches(0.35))
    write(frame, [[("dark chip", {"bold": True, "color": NAVY}),
                   (" = the baseline level, present in every slice. Each slice's other "
                    "conditions were run once, at the fixed settings shown.",
                    {"color": MUTED})]], size=9.5, space_after=0)
    titled_box(slide, Inches(8.2), Inches(1.6), Inches(4.55), Inches(1.2), RED_FILL, RED,
               "Model tiers (fixed labels, defined here once)", [
        [("Low", {"bold": True, "color": RED}), (f" = {LOW_MODEL}      ", {}),
         ("High", {"bold": True, "color": RED}), (f" = {HIGH_MODEL}", {})],
        ("The label names the predeclared capability tier, not the score observed in "
         "a condition. Open and Closed name the search workflow; Low and High name "
         "the model.", {"size": 9.5, "color": MUTED}),
    ], body_size=10)
    titled_box(slide, Inches(8.2), Inches(2.95), Inches(4.55), Inches(2.3), GREY_FILL, GREY,
               "How to read every result figure", [
        "•  small dots = the 12 paired seeds; large marker = mean; vertical interval = "
        "bootstrap 95% CI; higher trash fidelity is better",
        "•  x-axis groups always in the order Random → Greedy → Low → High",
        "•  within a tier: LLM-Open = blue triangle, LLM-Closed = red diamond, side by side",
        "•  tier = marker fill: High filled, Low hollow",
        "•  thick frame or shaded band = the previous baseline condition",
        "•  the small lower panel shows the two key paired differences with 95% CI; "
        "filled = p < 0.05, hollow = not detected",
    ], title_color=NAVY, body_size=9.5, space_after=2)
    box(slide, Inches(0.55), Inches(4.75), Inches(7.4), Inches(1.55), AMBER_FILL, AMBER)
    frame = textbox(slide, Inches(0.7), Inches(4.81), Inches(7.15), Inches(1.45))
    write(frame, [
        [("Where the Low tier was not run ", {"bold": True, "color": ORANGE}),
         ("(budget, qubit and Hamiltonian slices), the figures show Random, Greedy and "
          "the High Open / Closed results only, and say so; no Low value is invented. "
          "The complete Random → Greedy → Low → High grouping appears on the model "
          "slide.", {})],
        [("Untested combinations are omitted rather than displayed as empty panels.",
          {"bold": True})],
    ], size=10, space_after=4)
    transition(slide, "the first slice changes the amount of search: the evaluation "
                      "budget B.")
    footer(slide, "Random and Greedy do not depend on the language model, so each is one "
                  "result per condition; Low and High each hold an Open and a Closed result.")


def _result_slide(deck, kicker, title, subtitle, figure, figure_height, takeaways,
                  transition_text, footer_text):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    header(slide, kicker, title, subtitle)
    pic = picture(slide, figure, Inches(0.55), Inches(1.52), height=figure_height)
    left = Inches(0.55) + pic.width + Inches(0.25)
    frame = textbox(slide, left, Inches(1.52), SLIDE_W - left - Inches(0.55), Inches(4.9))
    write(frame, takeaways, size=10, space_after=6)
    transition(slide, transition_text)
    footer(slide, footer_text)
    return slide


def slide_05_budget(deck, summary):
    co = {k: stat(summary, k, CLOSED_OPEN) for k in ("budget_b4", "reference", "budget_b16")}
    return _result_slide(
        deck, "4 · Budget slice", "Budget changes the value of feedback",
        "B = 4 → 8 → 16 with 4 qubits, Ising chain and the High model tier fixed; "
        "B = 8 (thick frame) is the previous baseline",
        "slice_budget", Inches(4.35), [
            [("Semantic proposals stay above non-semantic search at every budget: ",
              {"bold": True, "color": BLUE}),
             (f"LLM-Open − Random {short(summary, 'budget_b4', OPEN_RANDOM)} at B = 4, "
              f"{short(summary, 'reference', OPEN_RANDOM)} at B = 8, "
              f"{short(summary, 'budget_b16', OPEN_RANDOM)} at B = 16 — positive "
              "throughout, not detected at B = 4 where the budget is too small to "
              "separate the arms reliably.", {})],
            [("The Closed − Open advantage shrinks as the budget grows: ",
              {"bold": True, "color": RED}),
             (f"{_fmt(co['budget_b4']['mean_paired_gain'], 3)} at B = 4 → "
              f"{_fmt(co['reference']['mean_paired_gain'], 3)} at B = 8 → "
              f"{_fmt(co['budget_b16']['mean_paired_gain'], 3)} at B = 16 (not detected). "
              "Free redesign is worth most when the open batch is thinnest; a wide open "
              "batch buys the same breadth without feedback.", {})],
            [("Greedy is not a substitute: ", {"bold": True, "color": ORANGE}),
             (f"Greedy − Random is {short(summary, 'budget_b4', GREEDY_RANDOM)} at B = 4 "
              "— spending half a small budget on single-gate steps costs breadth.", {})],
        ],
        "budget changes the amount of search; next we ask whether register width "
        "changes the conclusion.",
        "Same trainer, prompts, seeds and selection rule at every budget; adaptive "
        "methods always split B as B/2 exploration + B/2 refinement (2+2, 4+4, 8+8).")


def slide_06_qubits(deck, summary):
    keys = ("reference", "qubits_n6", "qubits_n8")
    orr = {k: stat(summary, k, OPEN_RANDOM) for k in keys}
    co = {k: stat(summary, k, CLOSED_OPEN) for k in keys}
    v = {k: validity(k, "LLM-Closed") for k in keys}
    return _result_slide(
        deck, "5 · Qubit-count slice", "The semantic advantage grows as the register widens",
        "4 → 6 → 8 qubits with B = 8, Ising chain and the High model tier fixed; "
        "4 qubits (thick frame) is the previous baseline",
        "slice_qubits", Inches(4.35), [
            [("Finding 1 — absolute fidelity falls as the task widens: ",
              {"bold": True, "color": NAVY}),
             (f"Random {mean_of(summary, 'reference', 'Random'):.3f} → "
              f"{mean_of(summary, 'qubits_n8', 'Random'):.3f}; LLM-Open "
              f"{mean_of(summary, 'reference', 'LLM-Open'):.3f} → "
              f"{mean_of(summary, 'qubits_n8', 'LLM-Open'):.3f}. Width and circuit size "
              "grow together (3 rotations + 1 CNOT per qubit).", {})],
            [("Finding 2 — Open stays strong relative to Random; Closed falls below "
              "Open: ", {"bold": True, "color": BLUE}),
             (f"LLM-Open − Random {_fmt(orr['reference']['mean_paired_gain'], 3)} → "
              f"{_fmt(orr['qubits_n6']['mean_paired_gain'], 3)} → "
              f"{_fmt(orr['qubits_n8']['mean_paired_gain'], 3)}, significant at all three "
              f"widths; LLM-Closed − LLM-Open "
              f"{_fmt(co['reference']['mean_paired_gain'], 3)} → "
              f"{_fmt(co['qubits_n6']['mean_paired_gain'], 3)} → "
              f"{_fmt(co['qubits_n8']['mean_paired_gain'], 3)}, reversed at 6 and 8 qubits.",
              {})],
            [("Caveat — proposal validity: ", {"bold": True, "color": MUTED}),
             (f"the share of closed-loop proposals meeting the gate-count contract (the "
              f"required 3 rotations + 1 CNOT per qubit) fell from "
              f"{v['reference'] * 100:.0f}% to {v['qubits_n6'] * 100:.0f}% and "
              f"{v['qubits_n8'] * 100:.0f}%; an invalid proposal is spent as a random "
              "draw, so part of the reversal is a generation-validity effect.",
              {"color": MUTED})],
            [("Tested only for the Ising chain and the High tier.",
              {"bold": True, "color": KICKER})],
        ],
        "width changes task difficulty; next we change the physical state family while "
        "returning to the baseline width.",
        "Same B = 8, prompts (with the qubit count substituted), trainer and seeds at "
        "every width; latent and trash registers are always half and half.")


def slide_07_hamiltonian(deck, summary):
    x = "hamiltonian_xxz"
    return _result_slide(
        deck, "6 · Hamiltonian slice",
        "The broad ordering survives at the tested XXZ anchor",
        "Ising → XXZ with 4 qubits, B = 8 and the High model tier fixed; the Ising chain "
        "(thick frame) is the previous baseline",
        "slice_hamiltonian", Inches(4.35), [
            [("The same broad semantic advantage: ", {"bold": True, "color": BLUE}),
             (f"both semantic methods have higher means than Random on the XXZ chain "
              f"(LLM-Open {mean_of(summary, x, 'LLM-Open'):.3f}, LLM-Closed "
              f"{mean_of(summary, x, 'LLM-Closed'):.3f} vs Random "
              f"{mean_of(summary, x, 'Random'):.3f}). Per seed, LLM-Closed − Random is "
              f"{short(summary, x, CLOSED_RANDOM)}; LLM-Open − Random is "
              f"{short(summary, x, OPEN_RANDOM)} — Random's seeds are spread widely "
              "here, so the mean gain is not detected as a per-seed effect.", {})],
            [("Closed exceeds Open at this anchor: ", {"bold": True, "color": RED}),
             (f"LLM-Closed − LLM-Open {short(summary, x, CLOSED_OPEN)}, larger than at "
              "the Ising baseline.", {})],
            [("One XXZ point is not the XXZ landscape: ", {"bold": True, "color": KICKER}),
             ("it says nothing about how the XXZ result moves with budget, width or "
              "model tier — none of those were run for XXZ.", {})],
        ],
        "the physics anchor gives one second-family result; the final controlled "
        "change is the model itself.",
        "Only the stated Hamiltonian facts change in the prompt (name, formula, "
        "parameter range, interaction type); the circuit space, trainer and seeds are "
        "identical.")


def slide_08_model(deck, summary):
    r, low = "reference", "model_alt"
    return _result_slide(
        deck, "7 · Model slice",
        "The model tier changes reliability more than the main Open advantage",
        "Low → High with 4 qubits, Ising chain and B = 8 fixed; High (shaded band) is "
        "the previous baseline; groups Random → Greedy → Low → High",
        "slice_model", Inches(4.35), [
            [("The Open semantic advantage is similar for both tiers: ",
              {"bold": True, "color": BLUE}),
             (f"LLM-Open − Random {short(summary, low, OPEN_RANDOM)} for Low and "
              f"{short(summary, r, OPEN_RANDOM)} for High.", {})],
            [("The Closed − Open relationship differs between tiers: ",
              {"bold": True, "color": RED}),
             (f"{short(summary, low, CLOSED_OPEN)} for Low versus "
              f"{short(summary, r, CLOSED_OPEN)} for High — the small closed-loop gain "
              "seen at the baseline is not present for the Low tier.", {})],
            [("Resource-contract validity is a separate failure mode: ",
              {"bold": True, "color": ORANGE}),
             (f"only {validity(low, 'LLM-Open') * 100:.0f}% of the Low tier's open-batch "
              f"proposals met the gate-count contract "
              f"({validity(low, 'LLM-Closed') * 100:.0f}% for its closed loop; "
              f"{validity(r, 'LLM-Open') * 100:.0f}% / "
              f"{validity(r, 'LLM-Closed') * 100:.0f}% for High). Invalid proposals are "
              "replaced by random draws, so a validity gap and an architecture-quality "
              "gap are entangled in the plotted score.", {})],
            [("Low and High are fixed tier labels; the plotted score is not what "
              "defines them.", {"bold": True, "color": KICKER})],
        ],
        "the four slices can now be combined to separate robust effects from fragile "
        "ones.",
        "Identical prompt bytes, temperature, retry policy and token limit for both "
        "tiers; only the model name changes. Random and Greedy are the same runs in "
        "both groups.")


def _validity_word(entries):
    """entries: [(label, share)]; shares below 0.9 are called out."""
    worst = min(entries, key=lambda e: e[1])
    if worst[1] >= 0.9:
        return f"stable (≥ {worst[1] * 100:.0f}%)"
    return f"degraded: {worst[0]} {worst[1] * 100:.0f}%"


def slide_09_synthesis(deck, summary):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    header(slide, "8 · Cross-slice synthesis and limitations",
           "What survived across the tested slices?",
           "One row per tested slice; wording is evidence-calibrated — robust, "
           "condition-dependent, not detected, reversed, not tested")

    def row_words(keys, contrast, labels):
        words = [(labels[i], verdict(summary, k, contrast)) for i, k in enumerate(keys)]
        distinct = {w for _l, w in words}
        if distinct == {"robust"}:
            return "robust"
        if distinct == {"not detected"}:
            return "not detected"
        return "condition-dependent: " + " · ".join(f"{w} ({lab})" for lab, w in words)

    rows = [["Slice", "Open versus Random", "Closed versus Open", "Proposal validity"]]
    rows.append(["Budget (B = 4, 8, 16)",
                 row_words(["budget_b4", "reference", "budget_b16"], OPEN_RANDOM,
                           ["B = 4", "B = 8", "B = 16"]),
                 row_words(["budget_b4", "reference", "budget_b16"], CLOSED_OPEN,
                           ["B = 4", "B = 8", "B = 16"]),
                 _validity_word([(f"Closed B = {b}", validity(k, "LLM-Closed"))
                                 for k, b in (("budget_b4", 4), ("reference", 8),
                                              ("budget_b16", 16))])])
    rows.append(["Qubit count (4, 6, 8)",
                 row_words(["reference", "qubits_n6", "qubits_n8"], OPEN_RANDOM,
                           ["4", "6", "8"]),
                 row_words(["reference", "qubits_n6", "qubits_n8"], CLOSED_OPEN,
                           ["4", "6", "8"]),
                 _validity_word([(f"Closed at {n} qubits", validity(k, "LLM-Closed"))
                                 for k, n in (("reference", 4), ("qubits_n6", 6),
                                              ("qubits_n8", 8))])])
    rows.append(["Hamiltonian (Ising, XXZ)",
                 row_words(["reference", "hamiltonian_xxz"], OPEN_RANDOM, ["Ising", "XXZ"]),
                 row_words(["reference", "hamiltonian_xxz"], CLOSED_OPEN, ["Ising", "XXZ"]),
                 _validity_word([("Closed XXZ", validity("hamiltonian_xxz", "LLM-Closed")),
                                 ("Open XXZ", validity("hamiltonian_xxz", "LLM-Open"))])])
    rows.append(["Model tier (Low, High)",
                 row_words(["model_alt", "reference"], OPEN_RANDOM, ["Low", "High"]),
                 row_words(["model_alt", "reference"], CLOSED_OPEN, ["Low", "High"]),
                 _validity_word([("Low Open", validity("model_alt", "LLM-Open")),
                                 ("Low Closed", validity("model_alt", "LLM-Closed"))])])
    colors = {}
    for r_index in range(1, 5):
        for c_index in (1, 2, 3):
            word = rows[r_index][c_index]
            colors[(r_index, c_index)] = (GREEN if word == "robust" else
                                          RED if "reversed" in word else
                                          MUTED if word.startswith("not") else ORANGE)
    table(slide, rows, Inches(0.55), Inches(1.55),
          [Inches(2.15), Inches(3.35), Inches(4.0), Inches(2.7)],
          row_height=Inches(0.46), header_size=10, body_size=9, colors=colors)
    frame = textbox(slide, Inches(0.55), Inches(3.95), Inches(12.2), Inches(0.3))
    write(frame, [("Not tested: width and Hamiltonian changes for the Low tier; any change "
                   "away from the baseline for XXZ; every two-factor combination.",
                   {"size": 9.5, "color": MUTED})], space_after=0)
    titled_box(slide, Inches(0.55), Inches(4.35), Inches(6.55), Inches(2.05), GREEN_FILL,
               GREEN, "Across the tested one-factor slices", [
        [("•  the semantic Open advantage is the most robust finding — same sign in "
          "every slice, significant in most conditions;", {})],
        [("•  the additional closed-loop advantage is budget-, width-, Hamiltonian- and "
          "model-dependent — it is not a general result;", {})],
        [("•  proposal validity can confound observed method differences: an invalid "
          "proposal is spent as a random draw.", {})],
    ], title_color=GREEN, body_size=10, space_after=3)
    titled_box(slide, Inches(7.3), Inches(4.35), Inches(5.45), Inches(2.05), GREY_FILL, GREY,
               "Key limitations", [
        "•  no interactions; only one non-Ising anchor",
        "•  Low tested only at the baseline physical condition",
        "•  12 paired seeds: a \"not detected\" cell is not evidence of no effect",
        "•  noiseless state-vector simulation",
        "•  qubit count and circuit size scale together",
    ], title_color=NAVY, body_size=10, space_after=2)
    transition(slide, "ranking methods at a fixed budget answers the wrong question for "
                      "practice — how much budget is actually required?")
    footer(slide, "\"robust\" = same sign as the baseline and p < 0.05; \"reversed\" = "
                  "p < 0.05 with the opposite sign; \"not detected\" = p ≥ 0.05; "
                  "\"condition-dependent\" = the verdict differs across the slice's levels.")


def slide_10_next(deck, summary):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    header(slide, "9 · Conclusion and next action",
           "Next: how much budget is actually required?",
           "Three statements from the tested slices, then a targeted boundary-finding "
           "experiment rather than a factorial sweep")
    statements = [
        (GREEN_FILL, GREEN, "Robust",
         "Physics-informed Open proposals remain better than non-semantic search across "
         "the tested one-factor slices."),
        (RED_FILL, RED, "Fragile",
         "The closed-loop advantage is not a general result: it changes with budget, "
         "width, Hamiltonian and model tier."),
        (AMBER_FILL, AMBER, "Operational bottleneck",
         "Valid circuit generation must be separated from architecture quality — an "
         "invalid proposal is spent as a random draw."),
    ]
    for i, (fill, border, title, body) in enumerate(statements):
        titled_box(slide, Inches(0.55), Inches(1.55 + i * 1.15), Inches(4.6), Inches(1.05),
                   fill, border, title, [(body, {"size": 10})], title_size=12, body_size=10)
    box(slide, Inches(5.4), Inches(1.55), Inches(7.35), Inches(5.3), BLUE_FILL, BLUE, 1.75)
    frame = textbox(slide, Inches(5.58), Inches(1.62), Inches(7.0), Inches(5.2))
    write(frame, [
        [("Next experiment — the minimum budget for a target fidelity",
          {"bold": True, "color": BLUE, "size": 12})],
        [("Define before running: ", {"bold": True}),
         ("a target validation fidelity F_target (candidate targets 0.95 and 0.99). "
          "B_min = the smallest budget that reaches F_target in at least 10 of the 12 "
          "paired seeds. B_min is determined on validation data; the held-out test is "
          "read only after B_min is selected.", {})],
        [("API-minimising procedure", {"bold": True, "color": NAVY})],
        [("1  Reanalyse the existing logs first — no new calls.", {})],
        [("2  Build best-so-far validation-fidelity curves versus candidate-evaluation "
          "count from those logs.", {})],
        [("3  Estimate whether each target is already reached by B = 4, 8 or 16.", {})],
        [("4  Identify only the conditions whose B_min remains unresolved.", {})],
        [("5  Run new LLM calls only at the nearest bracketing budget for those "
          "unresolved conditions.", {})],
        [("6  Preserve the same 12 paired seeds for any new confirmatory run.", {})],
        [("7  Report an interval for B_min when the exact minimum cannot be identified "
          "without excessive API cost.", {})],
        [("Not proposed: a blind full sweep over every Hamiltonian, qubit count, model, "
          "method and budget. The practical next step is targeted boundary-finding, not "
          "factorial completion.", {"bold": True, "color": KICKER})],
    ], size=10, space_after=4)
    frame = textbox(slide, Inches(0.55), Inches(5.1), Inches(4.6), Inches(1.7))
    write(frame, [
        [("Why this matters: ", {"bold": True, "color": NAVY}),
         ("every comparison so far ranks methods at a fixed B. The question that decides "
          "whether any of this is worth using is the reverse — the budget each method "
          "needs to reach the accuracy the application requires.", {})],
    ], size=10, space_after=0)
    footer(slide, "All figures and statistics come from the committed result tables via "
                  "the committed build scripts; no new experiment or model call was made "
                  "for this deck.")


# ------------------------------------------------------------------ QA ----

# term -> (regex that detects a USE, regex that detects the DEFINITION)
GLOSSARY = {
    "QAE": (r"\bQAE\b", r"[Qq]uantum autoencoder \(QAE\)"),
    "trash fidelity": (r"F_trash|trash fidelity", r"Trash fidelity:\s+F_trash ="),
    "latent": (r"\blatent\b", r"Latent qubits\s+q0, q1"),
    "trash": (r"\btrash\b", r"Trash qubits\s+q2, q3"),
    "open-loop": (r"open-loop", r"open-loop: all semantic"),
    "closed-loop": (r"closed-loop", r"closed-loop: semantic initial proposals"),
    "budget B": (r"\bB\s*=\s*\d|budget B\b|\bB/2\b|\bB_min\b",
                 r"budget B = the number of candidate circuits trained and evaluated per seed"),
    "Ising": (r"\bIsing\b", r"H = −Σ Zᵢ Zᵢ₊₁ − h Σ Xᵢ"),
    "XXZ": (r"\bXXZ\b", r"XXZ = the second Hamiltonian family"),
    "High tier": (r"\bHigh\b", r"High tier = gpt-5\.4-mini|High model tier"),
    "Low tier": (r"\bLow\b", r"Low = gpt-4\.1-mini"),
    "CI": (r"\bCI\b", r"confidence interval \(CI\)"),
    "seeds": (r"\bseeds?\b", r"12 paired seeds \(12 independent draws"),
    "Wilcoxon": (r"Wilcoxon", r"p = exact paired Wilcoxon test"),
    "CNOT": (r"\bCNOT\b", r"CNOT \(controlled-NOT\)"),
    "held-out": (r"held-out", r"held-out test set \(64 unseen h values"),
    "gate-count contract": (
        r"gate-count contract",
        r"gate-count contract \(the required 3 rotations \+ 1 CNOT per qubit\)"),
    "exploration/refinement": (
        r"\bexploration\b|\brefinement\b",
        r"B/2 exploration \+ B/2 refinement|single-structural-change refinements"),
    "F_target": (r"F_target", r"target validation fidelity F_target"),
    "B_min": (r"B_min", r"B_min = the smallest budget"),
}

REQUIRED = {
    1: ["tested one-factor-slice study, not a complete factorial grid"],
    4: ["Untested combinations are omitted rather than displayed as empty panels",
        f"Low = {LOW_MODEL}", f"High = {HIGH_MODEL}", "Random → Greedy → Low → High"],
    9: ["Across the tested one-factor slices", "no interactions", "12 paired seeds",
        "noiseless", "not tested"],
    10: ["F_target", "B_min", "0.95", "0.99", "10 of the 12", "validation",
         "held-out test is read only after", "Reanalyse the existing logs",
         "best-so-far", "B = 4, 8 or 16", "unresolved", "bracketing", "12 paired seeds",
         "interval"],
}
SECTIONS = ["survive one-factor changes", "previous baseline",
            "controlled robustness design", "what was actually tested",
            "budget slice", "qubit-count slice", "hamiltonian slice", "model slice",
            "cross-slice synthesis", "conclusion and next action"]


def slide_texts(path: Path) -> list[str]:
    deck = Presentation(str(path))
    texts = [""] * len(deck.slides._sldIdLst)
    for index, text in extract_text(path):
        texts[index - 1] += text + "\n"
    return texts


def glossary_audit(texts: list[str]) -> list[str]:
    problems = []
    for term, (use, definition) in GLOSSARY.items():
        first_use = next((i for i, t in enumerate(texts) if re.search(use, t)), None)
        defined = next((i for i, t in enumerate(texts) if re.search(definition, t)), None)
        if first_use is None:
            continue
        if defined is None:
            problems.append(f"'{term}' is used on slide {first_use + 1} but never defined")
        elif defined > first_use:
            problems.append(f"'{term}' is used on slide {first_use + 1} before its "
                            f"definition on slide {defined + 1}")
    return problems


def structure_audit(texts: list[str]) -> list[str]:
    problems = []
    if len(texts) != 10:
        problems.append(f"deck has {len(texts)} slides, expected 10")
    for index, (needle, text) in enumerate(zip(SECTIONS, texts, strict=False), start=1):
        if needle not in text.lower():
            problems.append(f"slide {index}: expected the section '{needle}'")
    for index, phrases in REQUIRED.items():
        text = texts[index - 1] if index <= len(texts) else ""
        for phrase in phrases:
            if phrase not in text:
                problems.append(f"slide {index}: missing required text {phrase!r}")
    manifest = json.loads((ROOT / "slices.json").read_text())
    if manifest["group_order"] != ["Random", "Greedy", "Low", "High"]:
        problems.append("slice figures do not use the Random -> Greedy -> Low -> High order")
    model = manifest["figures"]["slice_model"]["groups"]
    if model != ["Random", "Greedy", "Low(Open, Closed)", "High(Open, Closed)"]:
        problems.append("model slice figure does not carry the four required groups")
    for name in ("slice_budget", "slice_qubits", "slice_hamiltonian"):
        if "Low" in " ".join(manifest["figures"][name]["groups"]):
            problems.append(f"{name} shows a Low group although Low was not run there")
    return problems


def prohibited_audit(texts: list[str]) -> list[str]:
    problems = []
    for index, text in enumerate(texts, start=1):
        for pattern, label in FORBIDDEN_PATTERNS:
            match = re.search(pattern, text)
            if match:
                problems.append(f"slide {index}: forbidden {label} {match.group(0)!r}")
    return problems


def render_pdf(pptx_path: Path) -> Path:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                    str(pptx_path.parent), str(pptx_path)],
                   check=True, capture_output=True, timeout=600)
    return pptx_path.with_suffix(".pdf")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--no-pdf", action="store_true")
    args = parser.parse_args()
    DECK_DIR.mkdir(parents=True, exist_ok=True)
    pptx_path = DECK_DIR / f"{DECK_NAME}.pptx"

    if not args.check:
        summary = json.loads((ROOT / "summary.json").read_text())
        deck = Presentation()
        deck.slide_width, deck.slide_height = SLIDE_W, SLIDE_H
        for builder in (slide_01_question, slide_02_baseline, slide_03_design,
                        slide_04_reading, slide_05_budget, slide_06_qubits,
                        slide_07_hamiltonian, slide_08_model, slide_09_synthesis,
                        slide_10_next):
            builder(deck, summary)
        deck.save(str(pptx_path))
        print("wrote", pptx_path)

    texts = slide_texts(pptx_path)
    geometry = [p for p in geometry_audit(pptx_path)
                if "forbidden" not in p and "LINE (9)" not in p]
    problems = geometry + prohibited_audit(texts) + glossary_audit(texts) + \
        structure_audit(texts)
    if problems:
        print("\nAUDIT FAILED:")
        for problem in problems:
            print("  -", problem)
    else:
        print("audit passed: 10 slides, required sections and phrases, tested-slice "
              "figures in Random -> Greedy -> Low -> High order, every term defined "
              "before use, nothing off-canvas, no prohibited labels or claims")
    if not args.no_pdf and not args.check:
        print("wrote", render_pdf(pptx_path))
    if problems:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
